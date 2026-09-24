#!/usr/bin/env python3
"""X (Twitter) 博主商务合作评估报告生成器。

两个子命令：

  fetch    通过 X API v2 拉取某账号近 N 天的全部推文（需要 X_BEARER_TOKEN）
  analyze  读取推文数据（X API JSON / Apify / twitterapi.io JSON / CSV），
           输出近 1 年流量、趋势、内容结构，以及近 3 个月推广广告的期数、流量和互动率

只依赖 Python 3.9+ 标准库。
"""

from __future__ import annotations

import argparse
import csv
import html
import json
import os
import re
import statistics
import sys
import time
import urllib.parse
import urllib.request
from collections import Counter, defaultdict
from dataclasses import dataclass, field
from datetime import datetime, timedelta, timezone
from pathlib import Path

CST = timezone(timedelta(hours=8))  # 报告统一使用北京时间

# ---------------------------------------------------------------------------
# 分类与广告识别词表（可按需在 config.json 中覆盖）
# ---------------------------------------------------------------------------

DEFAULT_CATEGORIES: dict[str, list[str]] = {
    "BTC 链上数据": ["btc", "bitcoin", "比特币", "链上", "筹码", "持仓", "换手", "成本价", "长期持有者",
                   "短期持有者", "lth", "sth", "glassnode", "获利盘", "抛压", "支撑", "阻力", "urpd", "矿工"],
    "宏观 / 美联储 / 美股": ["美联储", "fed", "fomc", "降息", "加息", "cpi", "pce", "非农", "就业", "通胀",
                        "美债", "美元", "流动性", "美股", "纳斯达克", "标普", "鲍威尔", "关税", "衰退", "gdp",
                        "利率", "财政部", "日本央行", "日元"],
    "政策监管": ["sec", "clarity", "genius", "法案", "监管", "合规", "参议", "众议", "立法", "特朗普", "川普",
             "trump", "白宫", "执法", "cftc", "牌照"],
    "ETF / 机构资金": ["etf", "贝莱德", "blackrock", "ibit", "mstr", "strategy", "微策略", "机构", "灰度",
                   "净流入", "净流出", "储备"],
    "ETH / 山寨 / 项目": ["eth", "以太坊", "sol", "bnb", "defi", "公链", "空投", "代币", "meme", "base",
                     "wlfi", "usd1", "稳定币", "usdt", "usdc", "hyperliquid", "项目"],
    "交易所 / 平台": ["binance", "币安", "okx", "欧易", "bitget", "bybit", "coinbase", "gate", "htx", "火币",
                 "交易所"],
    "地缘政治 / 时事": ["伊朗", "以色列", "俄罗斯", "乌克兰", "战争", "中东", "石油", "原油", "总统", "停火",
                   "制裁"],
    "港股 / A 股 / 传统资产": ["港股", "打新", "a股", "黄金", "恒指", "券商", "中签", "上证", "茅台", "外汇",
                         "原油期货", "差价合约", "cfd"],
    "生活 / 移民 / 税务": ["移民", "税", "签证", "身份", "新加坡", "日本", "泰国", "马来", "新山", "胡志明",
                     "生活", "医疗", "旅行", "美食"],
}

# 出现即高度疑似广告
AD_STRONG = ["#ad", "#广告", "#sponsored", "#赞助", "#合作", "#pr", "广告", "赞助", "sponsored",
             "合作推广", "商务合作", "本推由", "本帖由", "邀请码", "推荐码", "注册链接", "专属链接",
             "专属福利", "返佣", "赠金", "体验金", "新人福利", "开户福利", "开户链接", "点击注册",
             "点击链接", "promo code", "referral", "invite code", "限时福利", "粉丝专属"]
# 单独出现时弱信号
AD_WEAK = ["福利", "活动", "奖池", "抽奖", "合作", "partner", "注册", "开户", "入金", "送", "报名",
           "优惠", "手续费", "折扣", "瓜分", "首发", "上线", "限时"]
# 常见投放品牌（交易所 / 券商 / 外汇 / 钱包 / 数据工具）
AD_BRANDS: dict[str, list[str]] = {
    "Binance": ["binance", "币安"], "OKX": ["okx", "欧易"], "Bitget": ["bitget"], "Bybit": ["bybit"],
    "Gate": ["gate.io", "gate.com", "@gate", "gate "], "HTX": ["htx", "火币"], "MEXC": ["mexc", "抹茶"],
    "KuCoin": ["kucoin"], "Coinbase": ["coinbase"], "BingX": ["bingx"], "WEEX": ["weex"],
    "Backpack": ["backpack"], "Hyperliquid": ["hyperliquid"], "Aster": ["aster"], "Pionex": ["pionex", "派网"],
    "CoinW": ["coinw"], "LBank": ["lbank"], "BitMart": ["bitmart"], "Toobit": ["toobit"],
    "BloFin": ["blofin"], "Bitunix": ["bitunix"], "OSL": ["osl"], "HashKey": ["hashkey"],
    "TMGM": ["tmgm"], "IC Markets": ["ic markets", "icmarkets"], "Exness": ["exness"], "XM": [" xm ", "xm.com"],
    "Pepperstone": ["pepperstone"], "Doo Prime": ["doo prime", "dooprime"], "EBC": ["ebc金融", "ebc financial"],
    "Mitrade": ["mitrade"], "富途 / moomoo": ["富途", "futu", "moomoo"], "老虎证券": ["老虎证券", "tiger brokers"],
    "长桥": ["长桥", "longbridge"], "盈透": ["盈透", "ibkr", "interactive brokers"],
    "OneKey": ["onekey"], "Ledger": ["ledger"], "imToken": ["imtoken"], "SafePal": ["safepal"],
    "Glassnode": ["glassnode"], "CryptoQuant": ["cryptoquant"], "Nansen": ["nansen"],
}
REFERRAL_URL = re.compile(r"(ref=|refcode|ref_code|referral|invite|inviteCode|/join/|/register|"
                          r"utm_|affiliate|aff=|/r/|channel=|clickid)", re.I)


# ---------------------------------------------------------------------------
# 数据模型与规范化
# ---------------------------------------------------------------------------

@dataclass
class Post:
    id: str
    created_at: datetime
    text: str
    views: int | None = None
    likes: int = 0
    retweets: int = 0
    replies: int = 0
    quotes: int = 0
    bookmarks: int = 0
    is_reply: bool = False
    is_retweet: bool = False
    is_quote: bool = False
    urls: list[str] = field(default_factory=list)
    url: str = ""
    # 仅博主后台导出才有
    link_clicks: int | None = None
    profile_visits: int | None = None
    new_follows: int | None = None
    self_ad: bool | None = None      # 数据里自带的"是否广告"标注
    self_brand: str = ""
    # 分析字段
    category: str = "其他"
    ad_score: int = 0
    ad_reasons: list[str] = field(default_factory=list)
    ad_brand: str = ""
    is_ad: bool = False

    @property
    def interactions(self) -> int:
        return self.likes + self.retweets + self.replies + self.quotes + self.bookmarks

    @property
    def er(self) -> float | None:
        return self.interactions / self.views if self.views else None

    @property
    def kind(self) -> str:
        if self.is_retweet:
            return "转推"
        if self.is_reply:
            return "回复"
        if self.is_quote:
            return "引用"
        return "原创"


def _first(d: dict, *keys, default=None):
    for k in keys:
        cur = d
        ok = True
        for part in k.split("."):
            if isinstance(cur, dict) and part in cur and cur[part] not in (None, ""):
                cur = cur[part]
            else:
                ok = False
                break
        if ok:
            return cur
    return default


def _int(v) -> int | None:
    if v is None or v == "":
        return None
    try:
        s = str(v).replace(",", "").strip()
        mult = 1
        if s[-1:].upper() == "K":
            mult, s = 1_000, s[:-1]
        elif s[-1:].upper() == "M":
            mult, s = 1_000_000, s[:-1]
        elif s.endswith("万"):
            mult, s = 10_000, s[:-1]
        return int(float(s) * mult)
    except (ValueError, IndexError):
        return None


def _bool(v) -> bool:
    if isinstance(v, bool):
        return v
    return str(v).strip().lower() in {"1", "true", "yes", "y", "是"}


def _date(v) -> datetime | None:
    if v is None or v == "":
        return None
    if isinstance(v, (int, float)):
        ts = v / 1000 if v > 1e12 else v
        return datetime.fromtimestamp(ts, tz=timezone.utc)
    s = str(v).strip()
    for fmt in ("%a %b %d %H:%M:%S %z %Y", "%Y-%m-%dT%H:%M:%S.%f%z", "%Y-%m-%dT%H:%M:%S%z",
                "%Y-%m-%d %H:%M:%S%z", "%Y-%m-%d %H:%M:%S", "%Y-%m-%d %H:%M", "%Y-%m-%d", "%Y/%m/%d %H:%M",
                "%Y/%m/%d", "%Y-%m-%d %H:%M %z", "%a, %b %d, %Y", "%b %d, %Y"):
        try:
            d = datetime.strptime(s.replace("Z", "+0000"), fmt)
            return d if d.tzinfo else d.replace(tzinfo=CST)
        except ValueError:
            continue
    try:
        d = datetime.fromisoformat(s.replace("Z", "+00:00"))
        return d if d.tzinfo else d.replace(tzinfo=CST)
    except ValueError:
        return None


# 表格类数据（模板 / X 后台导出 / 手工整理）的列名 → 统一字段名，匹配时忽略大小写、空格和括号里的备注
COLUMN_ALIASES = {
    "id": ["推文id", "帖子id", "post id", "tweet id", "id"],
    "created_at": ["发布时间", "时间", "date", "time", "created at", "createdat"],
    "text": ["正文", "内容", "post text", "tweet text", "text"],
    "url": ["帖子链接", "链接", "post link", "tweet permalink", "url"],
    "类型": ["类型", "帖子类型", "type"],
    "viewCount": ["浏览量", "曝光", "曝光量", "impressions", "views", "viewcount"],
    "likeCount": ["点赞", "喜欢", "likes", "likecount"],
    "retweetCount": ["转推", "转发", "reposts", "retweets", "retweetcount"],
    "replyCount": ["评论", "回复数", "replies", "replycount"],
    "quoteCount": ["引用", "quotes", "quotecount"],
    "bookmarkCount": ["收藏", "书签", "bookmarks", "bookmarkcount"],
    "link_clicks": ["链接点击", "url clicks", "link clicks"],
    "profile_visits": ["主页访问", "profile visits", "user profile clicks"],
    "new_follows": ["新增关注", "new follows", "follows"],
    "self_ad": ["是否广告", "广告", "is_ad", "sponsored"],
    "self_brand": ["合作品牌", "品牌", "brand"],
    "urls": ["外链", "推广链接", "links"],
}
_ALIAS_INDEX = {a: canon for canon, names in COLUMN_ALIASES.items() for a in names}


def canon_row(row: dict) -> dict:
    out = {}
    for k, v in row.items():
        if k is None:
            continue
        key = re.sub(r"[（(].*?[)）]", "", str(k)).strip().lower()
        canon = _ALIAS_INDEX.get(key) or _ALIAS_INDEX.get(key.replace(" ", "")) or str(k)
        if canon not in out or out[canon] in (None, ""):
            out[canon] = v
    kind = str(out.get("类型") or "")
    if kind:
        out.setdefault("isReply", "回复" in kind or "reply" in kind.lower())
        out.setdefault("isRetweet", "转推" in kind or "retweet" in kind.lower() or "repost" in kind.lower())
        out.setdefault("isQuote", "引用" in kind or "quote" in kind.lower())
    return out


def normalize(raw: dict, username: str) -> Post | None:
    """兼容 X API v2、Apify tweet scraper、twitterapi.io、X 后台导出以及数据模板。"""
    pm = raw.get("public_metrics") or {}
    tid = str(_first(raw, "id", "id_str", "tweet_id", "tweetId", "推文ID", default="")).strip()
    link = str(_first(raw, "url", "twitterUrl", "链接", default="") or "")
    if not re.fullmatch(r"\d+", tid):
        m = re.search(r"/status/(\d+)", link)
        tid = m.group(1) if m else ""
    created = _date(_first(raw, "created_at", "createdAt", "date", "time", "timestamp", "发布时间"))
    if not tid or not created:
        return None
    text = _first(raw, "note_tweet.text", "full_text", "fullText", "text", "content", "正文", default="") or ""

    refs = raw.get("referenced_tweets") or []
    ref_types = {r.get("type") for r in refs if isinstance(r, dict)}
    is_rt = _bool(_first(raw, "isRetweet", "is_retweet", "retweeted", default=False)) or "retweeted" in ref_types \
        or bool(raw.get("retweeted_tweet") or raw.get("retweeted_status")) or str(text).startswith("RT @")
    is_reply = _bool(_first(raw, "isReply", "is_reply", default=False)) or "replied_to" in ref_types \
        or bool(_first(raw, "in_reply_to_user_id", "inReplyToId", "in_reply_to_status_id_str"))
    is_quote = _bool(_first(raw, "isQuote", "is_quote", "is_quote_status", default=False)) or "quoted" in ref_types
    # 回复自己 = 长推串，按原创计
    reply_to = str(_first(raw, "in_reply_to_user_id", "inReplyToUserId", default="") or "")
    reply_name = str(_first(raw, "inReplyToUsername", "in_reply_to_screen_name", default="") or "")
    author_id = str(_first(raw, "author_id", "author.id", default="") or "")
    if is_reply and ((reply_to and reply_to == author_id) or reply_name.lower() == username.lower()):
        is_reply = False

    urls: list[str] = []
    for u in (_first(raw, "entities.urls", default=[]) or []):
        if isinstance(u, dict):
            urls.append(u.get("unwound_url") or u.get("expanded_url") or u.get("url") or "")
    extra = _first(raw, "urls", "links", "外链", default=None)
    if isinstance(extra, str):
        urls += [x for x in re.split(r"[\s,;|]+", extra) if x]
    elif isinstance(extra, list):
        urls += [x if isinstance(x, str) else x.get("expanded_url", "") for x in extra]

    return Post(
        id=tid,
        created_at=created.astimezone(CST),
        text=str(text),
        views=_int(_first(pm, "impression_count") if pm else None) or _int(
            _first(raw, "viewCount", "views", "view_count", "impressions", "浏览量", "曝光")),
        likes=_int(_first(pm, "like_count") if pm else _first(raw, "likeCount", "favorite_count", "likes",
                                                                  "点赞")) or 0,
        retweets=_int(_first(pm, "retweet_count") if pm else _first(raw, "retweetCount", "retweet_count",
                                                                        "retweets", "转推")) or 0,
        replies=_int(_first(pm, "reply_count") if pm else _first(raw, "replyCount", "reply_count", "replies",
                                                                     "评论")) or 0,
        quotes=_int(_first(pm, "quote_count") if pm else _first(raw, "quoteCount", "quote_count", "quotes",
                                                                   "引用")) or 0,
        bookmarks=_int(_first(pm, "bookmark_count") if pm else _first(raw, "bookmarkCount", "bookmark_count",
                                                                         "bookmarks", "收藏")) or 0,
        is_reply=is_reply, is_retweet=is_rt, is_quote=is_quote,
        urls=[u for u in urls if u],
        url=link or f"https://x.com/{username}/status/{tid}",
        link_clicks=_int(raw.get("link_clicks")),
        profile_visits=_int(raw.get("profile_visits")),
        new_follows=_int(raw.get("new_follows")),
        self_ad=_bool(raw["self_ad"]) if str(raw.get("self_ad") or "").strip() else None,
        self_brand=str(raw.get("self_brand") or "").strip(),
    )


def read_xlsx(path: Path) -> list[dict]:
    """读取数据模板的"推文明细"工作表（没有则读第一个表）。需要 openpyxl。"""
    try:
        import openpyxl
    except ImportError:
        raise SystemExit("读取 .xlsx 需要 openpyxl：pip install openpyxl（或把工作表另存为 CSV）")
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    ws = next((wb[n] for n in wb.sheetnames if "推文" in n or "post" in n.lower()), wb.worksheets[0])
    it = ws.iter_rows(values_only=True)
    header = [str(h).strip() if h is not None else "" for h in next(it, [])]
    return [dict(zip(header, r)) for r in it if any(c not in (None, "") for c in r)]


def load_posts(path: Path, username: str) -> list[Post]:
    rows: list[dict] = []
    if path.suffix.lower() == ".csv":
        with path.open(encoding="utf-8-sig", newline="") as f:
            rows = [canon_row(r) for r in csv.DictReader(f)]
    elif path.suffix.lower() in (".xlsx", ".xlsm"):
        rows = [canon_row(r) for r in read_xlsx(path)]
    else:
        text = path.read_text(encoding="utf-8").strip()
        if path.suffix.lower() == ".jsonl" or (text.startswith("{") and "\n{" in text):
            rows = [json.loads(line) for line in text.splitlines() if line.strip()]
        else:
            data = json.loads(text)
            if isinstance(data, dict):
                data = data.get("data") or data.get("tweets") or data.get("items") or []
            rows = data
    posts: dict[str, Post] = {}
    skipped = 0
    for r in rows:
        p = normalize(r, username)
        if p:
            posts[p.id] = p
        else:
            skipped += 1
    if skipped:
        print(f"[warn] 跳过 {skipped} 条缺少 id / 时间的记录", file=sys.stderr)
    return sorted(posts.values(), key=lambda p: p.created_at)


# ---------------------------------------------------------------------------
# 分类 & 广告识别
# ---------------------------------------------------------------------------

def classify(p: Post, categories: dict[str, list[str]]) -> str:
    t = p.text.lower()
    scores = {c: sum(t.count(k) for k in kws) for c, kws in categories.items()}
    best = max(scores.items(), key=lambda kv: kv[1])
    return best[0] if best[1] > 0 else "其他"


def detect_ad(p: Post) -> None:
    t = f" {p.text.lower()} "
    reasons: list[str] = []
    score = 0
    for k in AD_STRONG:
        if k in t:
            score += 3
            reasons.append(f"关键词:{k}")
    weak_hits = [k for k in AD_WEAK if k in t]
    if weak_hits:
        score += min(len(weak_hits), 2)
        reasons.append("弱信号:" + "/".join(weak_hits[:4]))
    brands = [b for b, kws in AD_BRANDS.items() if any(k in t or any(k.strip() in u.lower() for u in p.urls)
                                                        for k in kws)]
    ref_links = [u for u in p.urls if REFERRAL_URL.search(u)]
    if ref_links:
        score += 3
        reasons.append("带参/注册链接")
    if brands and (ref_links or p.urls):
        score += 2
        reasons.append("品牌+外链")
    elif brands and weak_hits:
        score += 1
    # 结尾署名式植入，如"一个 @Gate ，交易更多市场"：最后一段很短且提到品牌
    lines = [ln.strip() for ln in p.text.lower().splitlines() if ln.strip()]
    tail = lines[-1] if len(lines) > 1 else ""
    tail_brands = [b for b, kws in AD_BRANDS.items() if any(k.strip() in tail for k in kws)] if tail else []
    if tail_brands and len(tail) <= 40:
        score += 3
        reasons.append("结尾品牌署名" + ("(@提及)" if "@" in tail else ""))
        brands = tail_brands + [b for b in brands if b not in tail_brands]
    p.ad_brand = " / ".join(brands[:3])
    p.ad_score = score
    p.ad_reasons = reasons
    p.is_ad = score >= 3 and not p.is_retweet


def load_overrides(path: Path | None) -> dict[str, tuple[bool, str]]:
    """人工复核结果：CSV 列 id,is_ad[,brand]，is_ad 留空的行忽略。优先级高于自动识别。"""
    if not path:
        return {}
    out = {}
    with path.open(encoding="utf-8-sig", newline="") as f:
        for row in csv.DictReader(f):
            if str(row.get("id", "")).strip() and str(row.get("is_ad", "")).strip():
                out[row["id"].strip()] = (_bool(row["is_ad"]), (row.get("brand") or "").strip())
    return out


# ---------------------------------------------------------------------------
# 统计
# ---------------------------------------------------------------------------

def _mean(xs):
    xs = [x for x in xs if x is not None]
    return statistics.fmean(xs) if xs else None


def _median(xs):
    xs = [x for x in xs if x is not None]
    return statistics.median(xs) if xs else None


def summarize(ps: list[Post], followers: int | None = None) -> dict:
    views = [p.views for p in ps if p.views is not None]
    tot_v = sum(views)
    tot_i = sum(p.interactions for p in ps if p.views is not None)
    return {
        "posts": len(ps),
        "with_views": len(views),
        "total_views": tot_v,
        "avg_views": _mean(views),
        "median_views": _median(views),
        "p90_views": sorted(views)[int(len(views) * 0.9)] if len(views) >= 10 else None,
        "avg_likes": _mean([p.likes for p in ps]),
        "avg_retweets": _mean([p.retweets for p in ps]),
        "avg_replies": _mean([p.replies for p in ps]),
        "avg_quotes": _mean([p.quotes for p in ps]),
        "avg_bookmarks": _mean([p.bookmarks for p in ps]),
        "avg_interactions": _mean([p.interactions for p in ps]),
        "er_weighted": tot_i / tot_v if tot_v else None,          # 总互动 / 总浏览
        "er_median": _median([p.er for p in ps]),                  # 单帖互动率中位数
        "er_followers": (_mean([p.interactions for p in ps]) / followers) if followers and ps else None,
        "avg_link_clicks": _mean([p.link_clicks for p in ps]),
        "ctr": (sum(p.link_clicks for p in ps if p.link_clicks is not None and p.views) /
                sum(p.views for p in ps if p.link_clicks is not None and p.views))
        if any(p.link_clicks is not None and p.views for p in ps) else None,
        "avg_profile_visits": _mean([p.profile_visits for p in ps]),
        "avg_new_follows": _mean([p.new_follows for p in ps]),
    }


def linear_slope(ys: list[float]) -> float | None:
    pts = [(i, y) for i, y in enumerate(ys) if y is not None]
    if len(pts) < 3:
        return None
    mx = statistics.fmean(x for x, _ in pts)
    my = statistics.fmean(y for _, y in pts)
    den = sum((x - mx) ** 2 for x, _ in pts)
    return sum((x - mx) * (y - my) for x, y in pts) / den if den else None


def pct(a, b):
    return (a - b) / b if a is not None and b else None


def analyze(posts: list[Post], as_of: datetime, days: int, ad_days: int, followers: int | None,
            quote_usd: float | None, categories: dict[str, list[str]],
            overrides: dict[str, tuple[bool, str]] | None = None) -> dict:
    start = as_of - timedelta(days=days)
    ad_start = as_of - timedelta(days=ad_days)
    window = [p for p in posts if start <= p.created_at <= as_of]
    for p in window:
        p.category = classify(p, categories)
        detect_ad(p)
        if p.self_ad is not None:
            p.is_ad = p.self_ad
            p.ad_brand = p.self_brand or p.ad_brand
            p.ad_reasons.append("数据自带标注")
        if overrides and p.id in overrides:
            p.is_ad, brand = overrides[p.id]
            p.ad_brand = brand or p.ad_brand
            p.ad_reasons.append("人工复核")
    main = [p for p in window if not p.is_retweet and not p.is_reply]  # 主贴：原创 + 引用 + 自回复长推
    replies = [p for p in window if p.is_reply]
    rts = [p for p in window if p.is_retweet]

    # 月度
    months: dict[str, list[Post]] = defaultdict(list)
    for p in main:
        months[p.created_at.strftime("%Y-%m")].append(p)
    monthly = []
    for m in sorted(months):
        s = summarize(months[m], followers)
        s["month"] = m
        s["ads"] = sum(1 for p in months[m] if p.is_ad)
        s["all_posts"] = sum(1 for p in window if p.created_at.strftime("%Y-%m") == m)
        monthly.append(s)

    # 近 3 个月 vs 前 3 个月 / 近 90 天 vs 全年
    def span(a, b):
        return [p for p in main if as_of - timedelta(days=a) < p.created_at <= as_of - timedelta(days=b)]
    last90, prev90 = span(90, 0), span(180, 90)
    s_last, s_prev = summarize(last90, followers), summarize(prev90, followers)
    trend = {
        "last90": s_last, "prev90": s_prev,
        "views_median_chg": pct(s_last["median_views"], s_prev["median_views"]),
        "views_avg_chg": pct(s_last["avg_views"], s_prev["avg_views"]),
        "er_chg": pct(s_last["er_weighted"], s_prev["er_weighted"]),
        "posts_chg": pct(s_last["posts"], s_prev["posts"]),
        "monthly_median_slope": linear_slope([m["median_views"] for m in monthly]),
    }

    # 内容结构
    cats = []
    for c in list(categories) + ["其他"]:
        ps = [p for p in main if p.category == c]
        if ps:
            s = summarize(ps, followers)
            s["category"] = c
            s["share"] = len(ps) / len(main)
            cats.append(s)
    cats.sort(key=lambda s: -s["posts"])

    # 发帖时间习惯
    hours = Counter(p.created_at.hour for p in main)
    weekdays = Counter(p.created_at.weekday() for p in main)
    hour_views = {h: _median([p.views for p in main if p.created_at.hour == h]) for h in range(24)}

    # 广告：近 ad_days 天
    ad_window = [p for p in main if p.created_at >= ad_start]
    ads = [p for p in ad_window if p.is_ad]
    organic = [p for p in ad_window if not p.is_ad]
    s_ads, s_org = summarize(ads, followers), summarize(organic, followers)
    # 期数：同一品牌 3 天内的连续广告算同一期
    campaigns: list[dict] = []
    for p in sorted(ads, key=lambda p: p.created_at):
        brand = p.ad_brand or "未识别品牌"
        last = next((c for c in reversed(campaigns) if c["brand"] == brand), None)
        if last and (p.created_at - last["end"]).days <= 3:
            last["posts"].append(p)
            last["end"] = p.created_at
        else:
            campaigns.append({"brand": brand, "start": p.created_at, "end": p.created_at, "posts": [p]})
    brand_stats = []
    for b in sorted({c["brand"] for c in campaigns}):
        ps = [p for c in campaigns if c["brand"] == b for p in c["posts"]]
        s = summarize(ps, followers)
        s["brand"] = b
        s["campaigns"] = sum(1 for c in campaigns if c["brand"] == b)
        brand_stats.append(s)
    brand_stats.sort(key=lambda s: -s["posts"])

    cpm = None
    if quote_usd and s_ads["median_views"]:
        cpm = quote_usd / s_ads["median_views"] * 1000
    elif quote_usd and s_org["median_views"]:
        cpm = quote_usd / s_org["median_views"] * 1000

    return {
        "as_of": as_of, "start": start, "ad_start": ad_start, "days": days, "ad_days": ad_days,
        "followers": followers, "quote_usd": quote_usd, "cpm": cpm,
        "counts": {"all": len(window), "main": len(main), "replies": len(replies), "retweets": len(rts),
                   "per_day": len(window) / days, "main_per_day": len(main) / days},
        "overall": summarize(main, followers),
        "replies_summary": summarize(replies, followers),
        "monthly": monthly, "trend": trend, "categories": cats,
        "hours": hours, "weekdays": weekdays, "hour_views": hour_views,
        "top": sorted([p for p in main if p.views], key=lambda p: -p.views)[:10],
        "top_er": sorted([p for p in main if p.views and p.views >= 1000], key=lambda p: -(p.er or 0))[:10],
        "ads": ads, "ad_summary": s_ads, "organic_summary": s_org, "campaigns": campaigns,
        "brand_stats": brand_stats,
        "ad_candidates": sorted([p for p in ad_window if p.ad_score > 0], key=lambda p: -p.ad_score),
        "window": window,
    }


# ---------------------------------------------------------------------------
# 输出
# ---------------------------------------------------------------------------

def fn(x, pct_=False, d=1):
    if x is None:
        return "—"
    if pct_:
        return f"{x * 100:.{max(d, 2) if abs(x) < 0.1 else d}f}%"
    if isinstance(x, float) and x < 100:
        return f"{x:,.{d}f}"
    x = round(x)
    if abs(x) >= 10_000:
        return f"{x / 10_000:,.1f}万"
    return f"{x:,}"


def signed(x):
    if x is None:
        return "—"
    return f"{'+' if x >= 0 else ''}{x * 100:.1f}%"


def short(t: str, n=60):
    t = re.sub(r"https?://\S+", "", t).replace("\n", " ").strip()
    return t[:n] + ("…" if len(t) > n else "")


def verdict(r: dict) -> list[str]:
    o, t, a, g = r["overall"], r["trend"], r["ad_summary"], r["organic_summary"]
    out = []
    if o["posts"] and o["with_views"] < o["posts"] * 0.5:
        out.append(f"⚠ 数据缺少浏览量：{o['posts']} 条主贴中只有 {o['with_views']} 条带 viewCount，浏览相关指标不可信，"
                   "请重新导出包含浏览量的数据；以下互动率改用粉丝数口径。")
    mc = t["views_median_chg"]
    if mc is not None:
        word = "上升" if mc > 0.1 else "下滑" if mc < -0.1 else "基本持平"
        out.append(f"流量趋势：近 90 天主贴浏览中位数 {fn(t['last90']['median_views'])}，较前 90 天 {signed(mc)}，"
                   f"整体{word}。")
    if o["er_weighted"] is None and o["er_followers"] is not None:
        out.append(f"互动质量：单帖平均互动 {fn(o['avg_interactions'])}，粉丝互动率 {fn(o['er_followers'], True)}"
                   "（平均互动/粉丝数）。")
    if o["er_weighted"] is not None:
        lvl = "较高" if o["er_weighted"] >= 0.02 else "中等" if o["er_weighted"] >= 0.008 else "偏低"
        out.append(f"互动质量：全年加权互动率 {fn(o['er_weighted'], True)}（互动/浏览），在加密类中文 KOL 中属于{lvl}水平"
                   "（经验区间：<0.8% 偏低，0.8%–2% 中等，>2% 较高）。")
    if r["ads"]:
        ratio = (a["median_views"] / g["median_views"]) if a["median_views"] and g["median_views"] else None
        basis = "浏览中位数"
        if ratio is None and a["avg_interactions"] and g["avg_interactions"]:
            ratio, basis = a["avg_interactions"] / g["avg_interactions"], "平均互动"
        out.append(f"商业化程度：近 {r['ad_days']} 天识别到 {len(r['ads'])} 条推广帖、{len(r['campaigns'])} 期合作，"
                   f"覆盖 {len(r['brand_stats'])} 个品牌；推广帖{basis}为自然帖的 "
                   f"{fn(ratio, True) if ratio else '—'}。")
        if ratio is not None:
            out.append("广告承接：推广帖流量衰减小，粉丝对商业内容接受度好。" if ratio >= 0.7 else
                       "广告承接：推广帖流量明显低于自然帖，需要在合作形式上做内容化包装（如结合行情分析植入）。")
    else:
        out.append(f"商业化程度：近 {r['ad_days']} 天未识别到明显推广帖（请结合 ad_candidates.csv 人工复核）。")
    brands = {b.strip() for s in r["brand_stats"] for b in s["brand"].split("/")}
    fx = {"IC Markets", "Exness", "XM", "Pepperstone", "Doo Prime", "EBC", "Mitrade"} & brands
    if "TMGM" in brands:
        out.append("合作记录：近期已有 TMGM 相关推广，可调取历史投放数据做续约对比。")
    if fx:
        out.append(f"竞品风险：近期已与外汇/差价合约同业 {', '.join(sorted(fx))} 合作，谈判时需确认排他条款。")
    if r["cpm"]:
        out.append(f"报价参考：按报价 ${r['quote_usd']:,.0f} 与单帖浏览中位数估算，CPM ≈ ${r['cpm']:,.2f}。")
    return out


def write_csvs(r: dict, out: Path) -> None:
    with (out / "posts.csv").open("w", encoding="utf-8-sig", newline="") as f:
        w = csv.writer(f)
        w.writerow(["id", "时间(北京)", "类型", "分类", "浏览", "点赞", "转推", "评论", "引用", "收藏", "互动", "互动率",
                    "是否广告", "广告品牌", "广告分", "识别依据", "正文", "链接"])
        for p in r["window"]:
            w.writerow([p.id, p.created_at.strftime("%Y-%m-%d %H:%M"), p.kind, p.category, p.views, p.likes,
                        p.retweets, p.replies, p.quotes, p.bookmarks, p.interactions,
                        f"{p.er:.4f}" if p.er is not None else "", int(p.is_ad), p.ad_brand, p.ad_score,
                        ";".join(p.ad_reasons), p.text, p.url])
    with (out / "monthly.csv").open("w", encoding="utf-8-sig", newline="") as f:
        w = csv.writer(f)
        w.writerow(["月份", "总发帖(含回复转推)", "主贴数", "浏览均值", "浏览中位数", "总浏览", "平均互动", "加权互动率",
                    "广告帖"])
        for m in r["monthly"]:
            w.writerow([m["month"], m["all_posts"], m["posts"], round(m["avg_views"] or 0),
                        round(m["median_views"] or 0), m["total_views"], round(m["avg_interactions"] or 0, 1),
                        f"{m['er_weighted']:.4f}" if m["er_weighted"] else "", m["ads"]])
    with (out / "ad_candidates.csv").open("w", encoding="utf-8-sig", newline="") as f:
        w = csv.writer(f)
        w.writerow(["id", "is_ad", "brand", "自动判定", "广告分", "识别依据", "时间(北京)", "浏览", "互动", "互动率",
                    "正文", "链接"])
        for p in r["ad_candidates"]:
            w.writerow([p.id, "", p.ad_brand, int(p.is_ad), p.ad_score, ";".join(p.ad_reasons),
                        p.created_at.strftime("%Y-%m-%d %H:%M"), p.views, p.interactions,
                        f"{p.er:.4f}" if p.er is not None else "", p.text, p.url])


def render_md(r: dict, user: str) -> str:
    o, c, t = r["overall"], r["counts"], r["trend"]
    L = [f"# @{user} 推广合作评估报告", "",
         f"统计区间：{r['start']:%Y-%m-%d} ~ {r['as_of']:%Y-%m-%d}（北京时间，近 {r['days']} 天）；"
         f"广告分析区间：{r['ad_start']:%Y-%m-%d} 起（近 {r['ad_days']} 天）。", "",
         "## 一、核心结论", ""] + [f"- {v}" for v in verdict(r)] + ["",
         "## 二、账号基础数据（近 1 年）", "",
         "| 指标 | 数值 |", "|---|---|",
         f"| 粉丝数 | {fn(r['followers']) if r['followers'] else '未提供'} |",
         f"| 总发帖（含回复/转推） | {c['all']:,}（日均 {c['per_day']:.1f}） |",
         f"| 主贴（原创+引用） | {c['main']:,}（日均 {c['main_per_day']:.1f}） |",
         f"| 回复 / 转推 | {c['replies']:,} / {c['retweets']:,} |",
         f"| 主贴总浏览 | {fn(o['total_views'])} |",
         f"| 单帖平均浏览 | {fn(o['avg_views'])} |",
         f"| 单帖浏览中位数 | {fn(o['median_views'])} |",
         f"| 单帖浏览 P90 | {fn(o['p90_views'])} |",
         f"| 平均点赞 / 转推 / 评论 / 收藏 | {fn(o['avg_likes'])} / {fn(o['avg_retweets'])} / {fn(o['avg_replies'])} / {fn(o['avg_bookmarks'])} |",
         f"| 加权互动率（总互动/总浏览） | {fn(o['er_weighted'], True)} |",
         f"| 单帖互动率中位数 | {fn(o['er_median'], True)} |",
         f"| 粉丝互动率（平均互动/粉丝） | {fn(o['er_followers'], True)} |", "",
         "> 说明：平均浏览会被爆款拉高，报价谈判建议以**中位数**为准。", "",
         "## 三、流量趋势", "",
         "| 月份 | 主贴数 | 浏览均值 | 浏览中位数 | 平均互动 | 加权互动率 | 广告帖 |", "|---|---|---|---|---|---|---|"]
    for m in r["monthly"]:
        L.append(f"| {m['month']} | {m['posts']} | {fn(m['avg_views'])} | {fn(m['median_views'])} | "
                 f"{fn(m['avg_interactions'])} | {fn(m['er_weighted'], True)} | {m['ads']} |")
    L += ["", "| 对比 | 前 90 天 | 近 90 天 | 变化 |", "|---|---|---|---|",
          f"| 主贴数 | {t['prev90']['posts']} | {t['last90']['posts']} | {signed(t['posts_chg'])} |",
          f"| 浏览中位数 | {fn(t['prev90']['median_views'])} | {fn(t['last90']['median_views'])} | {signed(t['views_median_chg'])} |",
          f"| 浏览均值 | {fn(t['prev90']['avg_views'])} | {fn(t['last90']['avg_views'])} | {signed(t['views_avg_chg'])} |",
          f"| 加权互动率 | {fn(t['prev90']['er_weighted'], True)} | {fn(t['last90']['er_weighted'], True)} | {signed(t['er_chg'])} |",
          "", "## 四、内容结构", "",
          "| 内容方向 | 帖数 | 占比 | 浏览中位数 | 加权互动率 |", "|---|---|---|---|---|"]
    for s in r["categories"]:
        L.append(f"| {s['category']} | {s['posts']} | {fn(s['share'], True)} | {fn(s['median_views'])} | "
                 f"{fn(s['er_weighted'], True)} |")
    peak = sorted(r["hours"].items(), key=lambda kv: -kv[1])[:3]
    L += ["", f"发帖高峰时段（北京时间）：{'、'.join(f'{h}:00' for h, _ in peak)}。", "",
          "### 浏览量 Top 10", "", "| 时间 | 浏览 | 互动 | 互动率 | 内容 |", "|---|---|---|---|---|"]
    for p in r["top"]:
        L.append(f"| {p.created_at:%m-%d} | {fn(p.views)} | {fn(p.interactions)} | {fn(p.er, True)} | "
                 f"[{short(p.text, 40)}]({p.url}) |")
    a, g = r["ad_summary"], r["organic_summary"]
    L += ["", f"## 五、近 {r['ad_days']} 天推广广告分析", "",
          f"- 推广帖：**{a['posts']} 条**，合作期数：**{len(r['campaigns'])} 期**（同品牌 3 天内连续发布视为 1 期），"
          f"涉及品牌 {len(r['brand_stats'])} 个。",
          f"- 同期主贴总数 {a['posts'] + g['posts']}，广告占比 "
          f"{fn(a['posts'] / (a['posts'] + g['posts']) if a['posts'] + g['posts'] else None, True)}。", "",
          "| 对比 | 推广帖 | 自然帖 |", "|---|---|---|",
          f"| 帖数 | {a['posts']} | {g['posts']} |",
          f"| 浏览均值 | {fn(a['avg_views'])} | {fn(g['avg_views'])} |",
          f"| 浏览中位数 | {fn(a['median_views'])} | {fn(g['median_views'])} |",
          f"| 平均互动 | {fn(a['avg_interactions'])} | {fn(g['avg_interactions'])} |",
          f"| 加权互动率 | {fn(a['er_weighted'], True)} | {fn(g['er_weighted'], True)} |",
          f"| 单帖互动率中位数 | {fn(a['er_median'], True)} | {fn(g['er_median'], True)} |"]
    if a["avg_link_clicks"] is not None or g["avg_link_clicks"] is not None:
        L += [f"| 平均链接点击 | {fn(a['avg_link_clicks'])} | {fn(g['avg_link_clicks'])} |",
              f"| 链接点击率（点击/浏览） | {fn(a['ctr'], True)} | {fn(g['ctr'], True)} |"]
    if a["avg_profile_visits"] is not None or g["avg_profile_visits"] is not None:
        L += [f"| 平均主页访问 | {fn(a['avg_profile_visits'])} | {fn(g['avg_profile_visits'])} |",
              f"| 平均新增关注 | {fn(a['avg_new_follows'])} | {fn(g['avg_new_follows'])} |"]
    L += ["",
          "| 品牌 | 期数 | 帖数 | 浏览中位数 | 平均互动 | 加权互动率 |", "|---|---|---|---|---|---|"]
    for b in r["brand_stats"]:
        L.append(f"| {b['brand']} | {b['campaigns']} | {b['posts']} | {fn(b['median_views'])} | "
                 f"{fn(b['avg_interactions'])} | {fn(b['er_weighted'], True)} |")
    L += ["", "### 推广帖明细", "", "| 期 | 时间 | 品牌 | 浏览 | 点赞 | 转推 | 评论 | 互动率 | 内容 |",
          "|---|---|---|---|---|---|---|---|---|"]
    for i, cpg in enumerate(r["campaigns"], 1):
        for p in cpg["posts"]:
            L.append(f"| {i} | {p.created_at:%m-%d %H:%M} | {cpg['brand']} | {fn(p.views)} | {p.likes} | {p.retweets} | "
                     f"{p.replies} | {fn(p.er, True)} | [{short(p.text, 36)}]({p.url}) |")
    L += ["", "## 六、方法说明", "",
          "- 主贴 = 原创 + 引用 + 自回复长推；回复他人与纯转推不计入流量均值。",
          "- 互动 = 点赞 + 转推 + 评论 + 引用 + 收藏；加权互动率 = 总互动 / 总浏览。",
          "- 广告识别基于关键词、品牌词、带参注册链接打分（≥3 分判定），自动结果已导出到 `ad_candidates.csv`，"
          "填写 `is_ad` 列后用 `--ad-review` 重跑即可得到人工校准后的数据。",
          "- 浏览量为 X 平台公开的 impression / view 计数。", ""]
    return "\n".join(L)


# --- HTML（内联 SVG 图表，无外部依赖） ---------------------------------------

def svg_bar_line(labels, bars, line, bar_name, line_name, w=760, h=260):
    pad_l, pad_r, pad_t, pad_b = 56, 56, 16, 36
    iw, ih = w - pad_l - pad_r, h - pad_t - pad_b
    n = max(len(labels), 1)
    bmax = max([b or 0 for b in bars] + [1])
    lvals = [v for v in line if v is not None]
    lmax = max(lvals + [1e-9])
    bw = iw / n * 0.6
    parts = [f'<svg viewBox="0 0 {w} {h}" class="chart" role="img">']
    for i in range(5):
        y = pad_t + ih * i / 4
        parts.append(f'<line x1="{pad_l}" x2="{w - pad_r}" y1="{y:.1f}" y2="{y:.1f}" class="grid"/>')
        parts.append(f'<text x="{pad_l - 6}" y="{y + 4:.1f}" class="ax" text-anchor="end">{fn(bmax * (1 - i / 4))}</text>')
        parts.append(f'<text x="{w - pad_r + 6}" y="{y + 4:.1f}" class="ax">{lmax * (1 - i / 4) * 100:.1f}%</text>')
    pts = []
    for i, lab in enumerate(labels):
        cx = pad_l + iw * (i + 0.5) / n
        bh = ih * (bars[i] or 0) / bmax
        parts.append(f'<rect x="{cx - bw / 2:.1f}" y="{pad_t + ih - bh:.1f}" width="{bw:.1f}" height="{bh:.1f}" '
                     f'rx="3" class="bar"><title>{lab} {bar_name} {fn(bars[i])}</title></rect>')
        parts.append(f'<text x="{cx:.1f}" y="{h - 14}" class="ax" text-anchor="middle">{lab[2:]}</text>')
        if line[i] is not None:
            pts.append((cx, pad_t + ih - ih * line[i] / lmax, lab, line[i]))
    if pts:
        parts.append('<polyline class="line" points="' + " ".join(f"{x:.1f},{y:.1f}" for x, y, *_ in pts) + '"/>')
        for x, y, lab, v in pts:
            parts.append(f'<circle cx="{x:.1f}" cy="{y:.1f}" r="3.5" class="dot"><title>{lab} {line_name} '
                         f'{v * 100:.2f}%</title></circle>')
    parts.append("</svg>")
    return "".join(parts)


def svg_hbar(items, w=760, fmt=lambda v: str(v)):
    row = 28
    h = row * len(items) + 8
    vmax = max([v for _, v in items] + [1])
    parts = [f'<svg viewBox="0 0 {w} {h}" class="chart" role="img">']
    for i, (lab, v) in enumerate(items):
        y = 4 + i * row
        bw = (w - 260) * v / vmax
        parts.append(f'<text x="170" y="{y + 17}" class="lab" text-anchor="end">{html.escape(lab)}</text>'
                     f'<rect x="180" y="{y + 4}" width="{bw:.1f}" height="18" rx="3" class="bar"/>'
                     f'<text x="{186 + bw:.1f}" y="{y + 17}" class="ax">{fmt(v)}</text>')
    parts.append("</svg>")
    return "".join(parts)


def md_table_to_html(md: str) -> str:
    """把 render_md 的 markdown 转成简单 HTML（只覆盖本报告用到的语法）。"""
    out, in_table, in_list = [], False, False
    for line in md.splitlines():
        s = line.strip()
        esc = html.escape(s)
        esc = re.sub(r"\[([^\]]+)\]\((https?://[^)]+)\)", r'<a href="\2" target="_blank">\1</a>', esc)
        esc = re.sub(r"\*\*(.+?)\*\*", r"<b>\1</b>", esc)
        esc = re.sub(r"`(.+?)`", r"<code>\1</code>", esc)
        if s.startswith("|"):
            cells = [c.strip() for c in esc.strip("|").split("|")]
            if set(s.replace("|", "").strip()) <= {"-"}:
                continue
            if not in_table:
                out.append('<div class="tw"><table><thead><tr>' + "".join(f"<th>{c}</th>" for c in cells)
                           + "</tr></thead><tbody>")
                in_table = True
            else:
                out.append("<tr>" + "".join(f"<td>{c}</td>" for c in cells) + "</tr>")
            continue
        if in_table:
            out.append("</tbody></table></div>")
            in_table = False
        if s.startswith("- "):
            if not in_list:
                out.append("<ul>")
                in_list = True
            out.append(f"<li>{esc[2:]}</li>")
            continue
        if in_list:
            out.append("</ul>")
            in_list = False
        if s.startswith("### "):
            out.append(f"<h3>{esc[4:]}</h3>")
        elif s.startswith("## "):
            out.append(f"<h2>{esc[3:]}</h2>")
        elif s.startswith("# "):
            out.append(f"<h1>{esc[2:]}</h1>")
        elif s.startswith("> "):
            out.append(f'<p class="note">{esc[5:]}</p>')
        elif s:
            out.append(f"<p>{esc}</p>")
    if in_table:
        out.append("</tbody></table></div>")
    if in_list:
        out.append("</ul>")
    return "\n".join(out)


CSS = """
:root{--bg:#fbfaf7;--fg:#1f2328;--muted:#6b7280;--line:#e5e2da;--card:#fff;--accent:#2563eb;--accent2:#d97706}
@media (prefers-color-scheme:dark){:root{--bg:#15171a;--fg:#e8e6e1;--muted:#9aa0a6;--line:#2c2f34;--card:#1c1f23;--accent:#60a5fa;--accent2:#fbbf24}}
*{box-sizing:border-box}body{margin:0;background:var(--bg);color:var(--fg);font:15px/1.65 -apple-system,"PingFang SC","Microsoft YaHei",sans-serif}
main{max-width:960px;margin:0 auto;padding:32px 16px 64px}h1{font-size:26px;margin:0 0 8px}h2{font-size:19px;margin:36px 0 12px;padding-top:12px;border-top:1px solid var(--line)}
h3{font-size:16px;margin:24px 0 8px}p,li{color:var(--fg)}.note{color:var(--muted);font-size:13px}a{color:var(--accent)}
.kpis{display:grid;grid-template-columns:repeat(auto-fit,minmax(160px,1fr));gap:12px;margin:20px 0}
.kpi{background:var(--card);border:1px solid var(--line);border-radius:10px;padding:12px 14px}.kpi b{display:block;font-size:22px;font-variant-numeric:tabular-nums}.kpi span{color:var(--muted);font-size:12px}
.tw{overflow-x:auto;margin:8px 0 16px}table{border-collapse:collapse;width:100%;font-size:13px;font-variant-numeric:tabular-nums}
th,td{border-bottom:1px solid var(--line);padding:6px 8px;text-align:left;white-space:nowrap}td:last-child{white-space:normal;min-width:180px}th{color:var(--muted);font-weight:600}
.chart{width:100%;height:auto;background:var(--card);border:1px solid var(--line);border-radius:10px;margin:8px 0}
.grid{stroke:var(--line)}.ax{fill:var(--muted);font-size:11px}.lab{fill:var(--fg);font-size:12px}.bar{fill:var(--accent);opacity:.85}
.line{fill:none;stroke:var(--accent2);stroke-width:2.5}.dot{fill:var(--accent2)}.legend{color:var(--muted);font-size:12px}
.sw{display:inline-block;width:10px;height:10px;border-radius:2px;margin:0 4px 0 12px}code{font-size:12px}
"""


def render_html(r: dict, user: str, md: str) -> str:
    o, a = r["overall"], r["ad_summary"]
    kpis = [("粉丝数", fn(r["followers"]) if r["followers"] else "—"),
            ("近 1 年主贴", f"{r['counts']['main']:,}"),
            ("单帖浏览中位数", fn(o["median_views"])),
            ("单帖平均浏览", fn(o["avg_views"])),
            ("加权互动率", fn(o["er_weighted"], True)),
            ("近 90 天流量变化", signed(r["trend"]["views_median_chg"])),
            (f"近 {r['ad_days']} 天推广帖", f"{a['posts']} 条 / {len(r['campaigns'])} 期"),
            ("推广帖浏览中位数", fn(a["median_views"]))]
    kpi_html = '<div class="kpis">' + "".join(f'<div class="kpi"><span>{k}</span><b>{v}</b></div>'
                                              for k, v in kpis) + "</div>"
    ms = r["monthly"]
    chart1 = svg_bar_line([m["month"] for m in ms], [m["median_views"] for m in ms],
                          [m["er_weighted"] for m in ms], "浏览中位数", "互动率")
    legend = ('<div class="legend"><span class="sw" style="background:var(--accent)"></span>月度单帖浏览中位数'
              '<span class="sw" style="background:var(--accent2)"></span>月度加权互动率（右轴）</div>')
    chart2 = svg_hbar([(s["category"], s["posts"]) for s in r["categories"]], fmt=lambda v: f"{v} 帖")
    body = md_table_to_html(md)
    body = body.replace("<h2>二、", kpi_html + "<h2>二、", 1)
    body = body.replace("<h2>四、", "<h3>月度趋势</h3>" + legend + chart1 + "<h2>四、", 1)
    body = body.replace("<h2>四、内容结构</h2>", "<h2>四、内容结构</h2>" + chart2, 1)
    return (f'<!doctype html><html lang="zh-CN"><head><meta charset="utf-8">'
            f'<meta name="viewport" content="width=device-width,initial-scale=1">'
            f"<title>@{html.escape(user)} 合作评估</title><style>{CSS}</style></head><body><main>{body}</main></body></html>")


# ---------------------------------------------------------------------------
# X API 抓取
# ---------------------------------------------------------------------------

def _api(url: str, token: str) -> dict:
    req = urllib.request.Request(url, headers={"Authorization": f"Bearer {token}", "User-Agent": "x-report/1.0"})
    last: Exception | None = None
    for attempt in range(6):
        try:
            with urllib.request.urlopen(req, timeout=30) as resp:
                return json.loads(resp.read())
        except urllib.error.HTTPError as e:
            if e.code == 429:
                reset = int(e.headers.get("x-rate-limit-reset", time.time() + 60))
                wait = max(reset - time.time(), 5) + 1
                print(f"[rate-limit] 等待 {wait:.0f}s", file=sys.stderr)
                time.sleep(wait)
                continue
            raise SystemExit(f"X API 错误 {e.code}: {e.read().decode(errors='ignore')[:500]}")
        except urllib.error.URLError as e:
            time.sleep(2 ** attempt)
            last = e
    raise SystemExit(f"X API 请求失败: {last}")


def cmd_fetch(args) -> None:
    token = os.environ.get("X_BEARER_TOKEN")
    if not token:
        raise SystemExit("请先设置环境变量 X_BEARER_TOKEN（X 开发者后台 → Keys and tokens → Bearer Token）")
    base = "https://api.x.com/2"
    u = _api(f"{base}/users/by/username/{args.user}?user.fields=public_metrics,description,created_at,verified",
             token)
    if "data" not in u:
        raise SystemExit(f"找不到用户 {args.user}: {u}")
    user = u["data"]
    start = (datetime.now(timezone.utc) - timedelta(days=args.days)).strftime("%Y-%m-%dT%H:%M:%SZ")
    fields = "created_at,public_metrics,entities,referenced_tweets,in_reply_to_user_id,author_id,note_tweet,lang"
    params = {"max_results": 100, "start_time": start, "tweet.fields": fields}
    tweets, token_next = [], None
    while True:
        if token_next:
            params["pagination_token"] = token_next
        page = _api(f"{base}/users/{user['id']}/tweets?{urllib.parse.urlencode(params)}", token)
        tweets += page.get("data", [])
        token_next = page.get("meta", {}).get("next_token")
        print(f"已获取 {len(tweets)} 条", file=sys.stderr)
        if not token_next:
            break
    if len(tweets) >= 3150:
        print("[warn] 用户时间线接口最多返回最近约 3200 条；若需完整一年，请改用全量搜索接口或第三方导出。",
              file=sys.stderr)
    Path(args.out).write_text(json.dumps({"user": user, "data": tweets}, ensure_ascii=False, indent=1),
                              encoding="utf-8")
    print(f"已保存 {len(tweets)} 条到 {args.out}；粉丝数 {user['public_metrics']['followers_count']:,}")


def cmd_analyze(args) -> None:
    src = Path(args.input)
    posts = load_posts(src, args.user)
    if not posts:
        raise SystemExit("没有读到任何推文，请检查输入文件格式")
    followers = args.followers
    if followers is None and src.suffix.lower() == ".json":
        try:
            meta = json.loads(src.read_text(encoding="utf-8"))
            followers = meta.get("user", {}).get("public_metrics", {}).get("followers_count") if isinstance(meta, dict) else None
        except (json.JSONDecodeError, AttributeError):
            pass
    categories = DEFAULT_CATEGORIES
    if args.config:
        cfg = json.loads(Path(args.config).read_text(encoding="utf-8"))
        categories = cfg.get("categories", categories)
        AD_STRONG.extend(cfg.get("ad_strong", []))
        AD_WEAK.extend(cfg.get("ad_weak", []))
        AD_BRANDS.update(cfg.get("ad_brands", {}))
    as_of = (_date(args.as_of) if args.as_of else max(p.created_at for p in posts)).astimezone(CST)
    as_of = as_of.replace(hour=23, minute=59, second=59) if args.as_of else as_of

    overrides = load_overrides(Path(args.ad_review) if args.ad_review else None)
    if overrides:
        print(f"已加载 {len(overrides)} 条人工复核结果", file=sys.stderr)
    r = analyze(posts, as_of, args.days, args.ad_days, followers, args.quote_usd, categories, overrides)

    out = Path(args.out_dir)
    out.mkdir(parents=True, exist_ok=True)
    md = render_md(r, args.user)
    (out / "report.md").write_text(md, encoding="utf-8")
    (out / "report.html").write_text(render_html(r, args.user, md), encoding="utf-8")
    write_csvs(r, out)
    print(f"报告已生成：{out}/report.html, report.md, posts.csv, monthly.csv, ad_candidates.csv")
    for line in verdict(r):
        print("  · " + line)


def main(argv=None) -> None:
    ap = argparse.ArgumentParser(description="X 博主推广合作评估报告")
    sub = ap.add_subparsers(dest="cmd", required=True)

    f = sub.add_parser("fetch", help="通过 X API v2 拉取推文")
    f.add_argument("--user", required=True, help="用户名，不带 @")
    f.add_argument("--days", type=int, default=365)
    f.add_argument("--out", default="tweets.json")
    f.set_defaults(func=cmd_fetch)

    a = sub.add_parser("analyze", help="生成分析报告")
    a.add_argument("--input", required=True, help="推文数据：.json / .jsonl / .csv")
    a.add_argument("--user", required=True, help="用户名，不带 @")
    a.add_argument("--out-dir", default="report")
    a.add_argument("--days", type=int, default=365, help="整体分析天数")
    a.add_argument("--ad-days", type=int, default=90, help="广告分析天数")
    a.add_argument("--followers", type=int, help="粉丝数（X API 抓取时会自动读取）")
    a.add_argument("--as-of", help="截止日期 YYYY-MM-DD，默认取数据中最新一条")
    a.add_argument("--quote-usd", type=float, help="对方单条报价（美元），用于估算 CPM")
    a.add_argument("--ad-review", help="人工复核后的 ad_candidates.csv（填写 is_ad 列）")
    a.add_argument("--config", help="自定义分类/广告词表 JSON")
    a.set_defaults(func=cmd_analyze)

    args = ap.parse_args(argv)
    args.func(args)


if __name__ == "__main__":
    main()
