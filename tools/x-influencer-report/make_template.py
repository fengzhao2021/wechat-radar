#!/usr/bin/env python3
"""生成发给博主填写的数据模板：templates/KOL数据模板.xlsx（需要 openpyxl）。"""

from pathlib import Path

from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

OUT = Path(__file__).parent / "templates" / "KOL数据模板.xlsx"
HEAD = PatternFill("solid", fgColor="1F3A5F")
REQ = PatternFill("solid", fgColor="B45309")
HEAD_FONT = Font(bold=True, color="FFFFFF")
EXAMPLE_FONT = Font(italic=True, color="9CA3AF")
WRAP = Alignment(wrap_text=True, vertical="top")


def sheet(wb, title, cols, example=None, rows=0, validations=None, first=False):
    """cols: [(表头, 列宽, 是否必填, 批注)]"""
    ws = wb.active if first else wb.create_sheet()
    ws.title = title
    for i, (name, width, required, note) in enumerate(cols, 1):
        c = ws.cell(row=1, column=i, value=name)
        c.fill, c.font, c.alignment = (REQ if required else HEAD), HEAD_FONT, WRAP
        if note:
            c.comment = Comment(note, "TMGM BD")
        ws.column_dimensions[c.column_letter].width = width
    ws.row_dimensions[1].height = 34
    ws.freeze_panes = "A2"
    r = 2
    if example:
        for i, v in enumerate(example, 1):
            ws.cell(row=r, column=i, value=v).font = EXAMPLE_FONT
        r += 1
    for rr in rows or []:
        for i, v in enumerate(rr, 1):
            ws.cell(row=r, column=i, value=v).alignment = WRAP
        r += 1
    for col, options in (validations or {}).items():
        dv = DataValidation(type="list", formula1='"' + ",".join(options) + '"', allow_blank=True)
        ws.add_data_validation(dv)
        dv.add(f"{col}2:{col}5000")
    return ws


def main():
    wb = Workbook()

    guide = sheet(wb, "填写说明", [("项目", 18, False, None), ("说明", 100, False, None)], first=True, rows=[
        ("用途", "用于评估推广合作效果，数据仅用于内部评估，不会对外公开。"),
        ("时间范围", "「推文明细」：近 12 个月；「广告案例」：近 3 个月；其余为截至填写日的最新数据。"),
        ("最省事的方式", "在 X 网页版进入 Analytics（数据分析）→ 内容（Content）→ 选择近 1 年 → 导出数据（Export data），"
                     "直接把导出的 CSV 发给我们即可，不需要再抄进「推文明细」表。"
                     "导出文件自带曝光、互动、链接点击、主页访问、新增关注等字段。"),
        ("如果导出不了", "按「推文明细」表的列逐条填写，至少填写橙色表头的必填列；也可以只填近 3 个月。"),
        ("橙色表头", "必填；深蓝色表头为选填，能填尽量填，尤其是“链接点击”和“是否广告”。"),
        ("灰色斜体行", "示例，可删除。"),
        ("截图也可以", "「账号概况」「粉丝画像」可以直接提供 X 后台对应页面的截图（Analytics 总览页、受众 Audience 页）。"),
        ("粉丝画像", "受众地区分布对我们非常重要（涉及不同地区的合规要求），请尽量提供。"),
    ])
    for row in guide.iter_rows(min_row=2):
        row[0].font = Font(bold=True)

    sheet(wb, "推文明细", [
        ("推文ID", 22, False, "帖子链接末尾的数字；填了帖子链接可不填"),
        ("帖子链接", 44, True, "例如 https://x.com/PhyrexNi/status/2102852659174846850"),
        ("发布时间", 18, True, "格式：2026-09-23 20:08（北京时间）"),
        ("类型", 10, True, "原创 / 引用 / 回复 / 转推"),
        ("正文", 50, True, "全文，便于内容分类和广告识别"),
        ("浏览量(曝光)", 14, True, "X 显示的 Views / Impressions"),
        ("点赞", 9, True, None),
        ("转推", 9, True, None),
        ("评论", 9, True, None),
        ("引用", 9, False, None),
        ("收藏", 9, False, None),
        ("链接点击", 11, False, "后台 URL Clicks，衡量带链接推广帖的导流能力"),
        ("主页访问", 11, False, "后台 Profile visits"),
        ("新增关注", 11, False, "后台 New follows"),
        ("是否广告", 10, False, "是 / 否（商业合作帖请标“是”）"),
        ("合作品牌", 14, False, "广告帖填写品牌名"),
        ("外链", 30, False, "帖子里的链接，多个用空格分隔"),
    ], example=["示例", "https://x.com/PhyrexNi/status/示例", "2026-09-23 20:08", "原创",
                "今天晚上市场回调了一些……一个 @Gate ，交易更多市场", 35000, 26, 2, 8, 0, 1, 120, 45, 6, "是",
                "Gate", "https://www.gate.com/..."],
        validations={"D": ["原创", "引用", "回复", "转推"], "O": ["是", "否"]})

    sheet(wb, "账号概况", [("指标", 26, True, None), ("数值", 18, True, None), ("说明", 60, False, None)], rows=[
        ("X 账号", "@PhyrexNi", "如有多个账号（如 @Phyrex_Ni）请分别说明用途"),
        ("当前粉丝数", "", ""),
        ("12 个月前粉丝数（约）", "", "用于计算粉丝增长"),
        ("近 28 天总曝光", "", "Analytics 总览页"),
        ("近 90 天总曝光", "", ""),
        ("近 90 天互动率", "", "Analytics 总览页 Engagement rate"),
        ("近 90 天主页访问", "", ""),
        ("近 90 天新增关注", "", ""),
        ("是否认证（蓝 V / 金标）", "", ""),
        ("其他平台及粉丝数", "", "Telegram 频道 / 币安广场 / YouTube / 公众号等，可同步分发的渠道"),
    ])

    sheet(wb, "粉丝画像", [("维度", 14, True, None), ("分类", 22, True, None), ("占比", 10, True, "如 35%"),
                        ("说明", 40, False, None)], rows=[
        ("地区", "（Top 10 逐行填写）", "", "X 后台 Audience → Location"),
        ("地区", "", "", ""), ("地区", "", "", ""), ("地区", "", "", ""), ("地区", "", "", ""),
        ("语言", "中文", "", ""), ("语言", "英文", "", ""),
        ("性别", "男", "", ""), ("性别", "女", "", ""),
        ("年龄", "18-24", "", ""), ("年龄", "25-34", "", ""), ("年龄", "35-44", "", ""), ("年龄", "45+", "", ""),
        ("设备", "iOS", "", ""), ("设备", "Android", "", ""), ("设备", "网页", "", ""),
    ])

    sheet(wb, "广告案例(近3个月)", [
        ("发布日期", 13, True, None),
        ("合作品牌", 14, True, None),
        ("品牌类型", 14, True, "交易所 / 外汇·差价合约 / 券商 / 钱包·工具 / 项目方 / 其他"),
        ("帖子链接", 40, True, None),
        ("合作形式", 16, True, "单条推文 / 结尾署名植入 / 置顶 / 长文 / Space 直播 / 转推 / 系列套餐"),
        ("是否带注册链接", 12, False, None),
        ("浏览量", 11, True, None),
        ("链接点击", 11, False, None),
        ("注册数", 10, False, "若品牌方有反馈"),
        ("开户/入金数", 12, False, "若品牌方有反馈"),
        ("是否排他", 10, False, "该合作是否限制同类品牌"),
        ("合作期限", 14, False, "如 2026-07 ~ 2026-09"),
        ("备注", 30, False, None),
    ], validations={"C": ["交易所", "外汇·差价合约", "券商", "钱包·工具", "项目方", "其他"],
                    "E": ["单条推文", "结尾署名植入", "置顶", "长文", "Space 直播", "转推", "系列套餐"],
                    "F": ["是", "否"], "K": ["是", "否"]})

    sheet(wb, "报价与档期", [
        ("合作形式", 18, True, None), ("单价(USD)", 12, True, None), ("包含内容", 36, False, "如：1 条推文 + 置顶 24 小时"),
        ("可否带追踪链接", 14, False, None), ("是否接受外汇/差价合约品牌", 18, False, None),
        ("当前排他限制", 26, False, "是否与其他交易平台有排他协议"), ("最早档期", 12, False, None), ("备注", 30, False, None),
    ], rows=[("单条推文",), ("结尾署名植入（按月）",), ("置顶",), ("长文 / 深度分析",), ("Space 直播",), ("月度套餐",)],
        validations={"D": ["是", "否"], "E": ["是", "否"]})

    OUT.parent.mkdir(exist_ok=True)
    wb.save(OUT)
    print(f"已生成 {OUT}")


if __name__ == "__main__":
    main()
